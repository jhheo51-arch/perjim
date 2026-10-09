"""Add the final-card response evidence to a versioned submission workbook."""

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/archive/Trend-to-Trust-Signal-Checks-2026-10-09-previous.xlsx"
RESULT = ROOT / "docs/validation/final-card-check-results-2026-10-09-v01.json"
OUTPUT = ROOT / "docs/submission/Trend-to-Trust-Validation-Evidence-2026-10-09.xlsx"

BLUE = "0064FF"
INK = "191F28"
MID = "4E5968"
LIGHT = "F5F8FC"
PALE = "E8F3FF"
WHITE = "FFFFFF"


def style_base(sheet, max_row, max_col):
    sheet.sheet_view.showGridLines = False
    for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=10, color=INK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)


def set_widths(sheet, widths):
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


workbook = load_workbook(SOURCE)
for name in ["완성 카드 응답", "반응과 수정 판단"]:
    if name in workbook.sheetnames:
        del workbook[name]

for sheet in workbook.worksheets:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and not cell.value.startswith("="):
                for mark in (chr(183), "•", "‧", "ㆍ"):
                    cell.value = cell.value.replace(mark, ", ")

data = json.loads(RESULT.read_text(encoding="utf-8"))
brand_labels = {
    "EASY_UNDERSTANDING": "돈을 쉽게 이해하도록 돕는 브랜드",
    "SPENDING_JUDGE": "소비를 평가하고 줄이라고 말하는 브랜드",
    "AUTO_ANALYSIS": "생활을 자동으로 추론하는 기술 브랜드",
    "UNSURE": "잘 모르겠음",
}

raw = workbook.create_sheet("완성 카드 응답")
style_base(raw, 14, 9)
raw.merge_cells("A2:I2")
raw["A2"] = "완성 카드 첫 3장, 익명 응답 5건"
raw["A2"].font = Font(name="맑은 고딕", size=18, bold=True, color=BLUE)
raw.row_dimensions[2].height = 38
raw.merge_cells("A4:I4")
raw["A4"] = "2026-10-09 확인. 계산 근거, 가상 사례 구분, 기능 오해와 브랜드 역할을 확인했습니다. 광고 성과나 전체 고객 의견이 아닙니다."
raw["A4"].font = Font(name="맑은 고딕", size=10, color=MID)
raw.row_dimensions[4].height = 42
headers = ["참여자", "계산 근거", "가상 사례", "기능 오해", "브랜드 역할", "이해도", "계속 볼 의향", "기억에 남은 내용", "바꾸면 좋을 부분"]
for column, value in enumerate(headers, 1):
    cell = raw.cell(6, column, value)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="맑은 고딕", size=10, bold=True, color=WHITE)
raw.row_dimensions[6].height = 36
for row_index, item in enumerate(data["responses"], 7):
    values = [
        item["participantId"],
        "정답" if item["amountBasis"] == "FOUR_X_18000" else "오답 또는 불확실",
        "가상 사례로 구분" if item["storyType"] == "FICTIONAL" else "오답 또는 불확실",
        "네" if item["featureConfusion"] == "YES" else "아니요" if item["featureConfusion"] == "NO" else "잘 모르겠음",
        brand_labels.get(item["brandRole"], item["brandRole"]),
        item["clarity"],
        item["interest"],
        item.get("impression") or "입력 없음",
        item.get("confusing") or "입력 없음",
    ]
    for column, value in enumerate(values, 1):
        raw.cell(row_index, column, value)
        if row_index % 2 == 1:
            raw.cell(row_index, column).fill = PatternFill("solid", fgColor=LIGHT)
    raw.row_dimensions[row_index].height = 76
raw.merge_cells("A13:I13")
raw["A13"] = "원자료: docs/validation/final-card-check-results-2026-10-09-v01.json"
raw["A13"].font = Font(name="맑은 고딕", size=9, color=MID)
raw.row_dimensions[13].height = 28
set_widths(raw, {"A": 12, "B": 18, "C": 20, "D": 18, "E": 38, "F": 11, "G": 14, "H": 34, "I": 34})
raw.freeze_panes = "B7"

summary = workbook.create_sheet("반응과 수정 판단")
style_base(summary, 20, 4)
summary.merge_cells("A2:D2")
summary["A2"] = "Trend to Trust, 반응에서 다음 카드 수정까지"
summary["A2"].font = Font(name="맑은 고딕", size=18, bold=True, color=BLUE)
summary.row_dimensions[2].height = 38
summary.merge_cells("A4:D4")
summary["A4"] = "문장 검증 2회와 완성 카드 반응 5건을 콘텐츠 수정으로 연결했습니다. 작은 표본이며 실제 광고와 제휴는 진행하지 않았습니다."
summary["A4"].font = Font(name="맑은 고딕", size=10, color=MID)
summary.row_dimensions[4].height = 42
for column, value in enumerate(["단계", "확인한 값", "결과", "편집 판단"], 1):
    cell = summary.cell(6, column, value)
    cell.fill = PatternFill("solid", fgColor=INK)
    cell.font = Font(name="맑은 고딕", size=10, bold=True, color=WHITE)
summary.row_dimensions[6].height = 36
summary_rows = [
    ["완성 카드", "응답 수", 5, "첫 3장을 별도 설명 없이 제시"],
    ["완성 카드", "계산 근거 이해", 5, "72,000원의 계산식 유지"],
    ["완성 카드", "가상 사례 구분", 5, "가상 설정 표시는 전달됨"],
    ["완성 카드", "자동 분석 기능 오해", 4, "질문형 제목 제거"],
    ["완성 카드", "기능 오해 불확실", 1, "인물의 직접 기록을 제목에 표시"],
    ["완성 카드", "목표 브랜드 역할", 0, "브랜드 역할 전달이 부족했음을 기록"],
    ["완성 카드", "이해도 평균", 4.8, "이해도와 브랜드 역할을 분리해서 판단"],
    ["완성 카드", "계속 볼 의향 평균", 3.0, "콘텐츠 흥미를 성과로 과장하지 않음"],
]
for row_index, values in enumerate(summary_rows, 7):
    for column, value in enumerate(values, 1):
        summary.cell(row_index, column, value)
        if row_index % 2 == 1:
            summary.cell(row_index, column).fill = PatternFill("solid", fgColor=LIGHT)
    summary.row_dimensions[row_index].height = 52
summary["C13"].number_format = "0.0"
summary["C14"].number_format = "0.0"
summary.merge_cells("A16:D16")
summary["A16"] = "수정 전: 그날은 왜 주문했을까요?  →  수정 후: 그날의 이유는 직접 적었습니다."
summary["A16"].fill = PatternFill("solid", fgColor=PALE)
summary["A16"].font = Font(name="맑은 고딕", size=10, bold=True, color=BLUE)
summary.row_dimensions[16].height = 42
summary.merge_cells("A18:D18")
summary["A18"] = "한계: 응답 수는 5건입니다. 수정 카드 재확인, 실제 게시 반응, 광고 성과와 제휴 결과는 아직 없습니다."
summary["A18"].font = Font(name="맑은 고딕", size=9, color=MID)
summary.row_dimensions[18].height = 36
summary.merge_cells("A20:D20")
summary["A20"] = "공개 결과: https://jhheo51-arch.github.io/trend-to-trust/final-card-check.html"
summary["A20"].font = Font(name="맑은 고딕", size=9, color=BLUE)
summary["A20"].hyperlink = "https://jhheo51-arch.github.io/trend-to-trust/final-card-check.html"
summary.row_dimensions[20].height = 28
set_widths(summary, {"A": 18, "B": 28, "C": 18, "D": 66})
summary.freeze_panes = "A7"

workbook._sheets.remove(summary)
workbook._sheets.remove(raw)
workbook._sheets.insert(0, summary)
workbook._sheets.insert(1, raw)
workbook.active = 0

workbook.calculation.fullCalcOnLoad = True
workbook.calculation.forceFullCalc = True
workbook.calculation.calcMode = "auto"
workbook.save(OUTPUT)

check = load_workbook(OUTPUT, data_only=False)
assert "완성 카드 응답" in check.sheetnames
assert "반응과 수정 판단" in check.sheetnames
assert check["반응과 수정 판단"]["C7"].value == 5
assert check["반응과 수정 판단"]["C10"].value == 4
assert check["반응과 수정 판단"]["C12"].value == 0
assert check["완성 카드 응답"].max_row >= 13
print(OUTPUT)
